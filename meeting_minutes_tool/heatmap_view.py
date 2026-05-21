# -*- coding: utf-8 -*-
"""
热力图视图模块
提供会议纪要热力图界面，按 1-50 编号展示纪要分布，支持点击打开文件
"""

import json
import logging
import os
import re
import subprocess
import sys
import tkinter as tk
from tkinter import ttk, messagebox
from datetime import datetime
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

# 样式常量
CELL_SIZE = 72           # 每个单元格大小（像素）
COLS_PER_ROW = 10        # 每行显示的格子数
NUM_ROWS = 5             # 行数（共 50 格）
GAP = 3                  # 单元格间距
MEETING_COLORS = {
    "党委会": "#4A90D9",  # 蓝色系
    "董事会": "#E67E22",  # 橙色系
    "总经办": "#2ECC71",  # 绿色系
}
GRAY_COLOR = "#CCCCCC"   # 无纪要颜色
TEXT_COLOR = "#FFFFFF"   # 文字颜色
GRAY_TEXT = "#999999"    # 灰色文字


class HeatmapView(tk.Frame):
    """热力图视图"""

    def __init__(self, parent, config_callback):
        """
        :param parent: 父容器
        :param config_callback: 获取配置的回调函数
        """
        super().__init__(parent)
        self.config_callback = config_callback
        self.meeting_types = ["党委会", "董事会", "总经办"]
        self.current_year = str(datetime.now().year)
        self.current_type = self.meeting_types[0]

        # 缓存数据: {序号: (source_file, meeting_date, pdf_path)}
        self.meeting_data = {}

        # 可用年份列表
        self.available_years = self._get_available_years()

        self._create_widgets()
        self._load_data()

    def _get_available_years(self) -> list:
        """获取可用的年份列表（从当前年份往前推10年）"""
        current = datetime.now().year
        return [str(y) for y in range(current, current - 10, -1)]

    def _create_widgets(self):
        """创建界面组件"""
        # ========== 顶部控制栏 ==========
        control_frame = tk.Frame(self, bg="#F5F5F5", padx=15, pady=10)
        control_frame.pack(fill=tk.X)

        # 会议类型下拉
        tk.Label(control_frame, text="会议类型：", bg="#F5F5F5",
                 font=("微软雅黑", 11)).pack(side=tk.LEFT, padx=(0, 5))
        self.type_combo = ttk.Combobox(
            control_frame, values=self.meeting_types,
            state="readonly", width=10, font=("微软雅黑", 11)
        )
        self.type_combo.set(self.current_type)
        self.type_combo.pack(side=tk.LEFT, padx=(0, 20))
        self.type_combo.bind("<<ComboboxSelected>>", self._on_filter_change)

        # 年份下拉
        tk.Label(control_frame, text="年份：", bg="#F5F5F5",
                 font=("微软雅黑", 11)).pack(side=tk.LEFT, padx=(0, 5))
        self.year_combo = ttk.Combobox(
            control_frame, values=self.available_years,
            state="readonly", width=8, font=("微软雅黑", 11)
        )
        self.year_combo.set(self.current_year)
        self.year_combo.pack(side=tk.LEFT, padx=(0, 20))
        self.year_combo.bind("<<ComboboxSelected>>", self._on_filter_change)

        # 刷新按钮
        self.refresh_btn = tk.Button(
            control_frame, text="🔄 刷新", command=self._load_data,
            bg="#3498DB", fg="white", font=("微软雅黑", 10, "bold"),
            padx=12, pady=2, cursor="hand2"
        )
        self.refresh_btn.pack(side=tk.LEFT, padx=(0, 10))

        # 提示标签
        self.info_label = tk.Label(
            control_frame, text="", bg="#F5F5F5",
            font=("微软雅黑", 10), fg="#666666"
        )
        self.info_label.pack(side=tk.RIGHT, padx=(10, 0))

        # 分隔线
        separator = tk.Frame(self, height=2, bg="#DDDDDD")
        separator.pack(fill=tk.X)

        # ========== 热力图画布区域 ==========
        canvas_frame = tk.Frame(self, bg="white")
        canvas_frame.pack(fill=tk.BOTH, expand=True)

        self.canvas = tk.Canvas(canvas_frame, bg="white", highlightthickness=0)
        self.scrollbar = tk.Scrollbar(
            canvas_frame, orient=tk.VERTICAL, command=self.canvas.yview
        )
        self.scrollable_frame = tk.Frame(self.canvas, bg="white")

        self.scrollable_frame.bind(
            "<Configure>",
            lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        )

        self.canvas_window = self.canvas.create_window(
            (0, 0), window=self.scrollable_frame, anchor="nw"
        )
        self.canvas.configure(yscrollcommand=self.scrollbar.set)

        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        # 绑定鼠标滚轮
        self._bind_mousewheel()

        # 绑定画布大小变化
        self.canvas.bind("<Configure>", self._on_canvas_configure)

    def _bind_mousewheel(self):
        """绑定鼠标滚轮事件"""
        def _on_mousewheel(event):
            self.canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
        self.canvas.bind_all("<MouseWheel>", _on_mousewheel)

    def _on_canvas_configure(self, event):
        """画布大小变化时调整内部 frame 宽度"""
        self.canvas.itemconfig(self.canvas_window, width=event.width)

    def _on_filter_change(self, event=None):
        """筛选条件变化"""
        self.current_type = self.type_combo.get()
        self.current_year = self.year_combo.get()
        self._load_data()

    def _get_config(self) -> dict:
        """获取当前配置"""
        try:
            return self.config_callback()
        except Exception:
            return {}

    def _load_data(self):
        """加载数据并渲染热力图"""
        self.current_type = self.type_combo.get()
        self.current_year = self.year_combo.get()
        self.meeting_data = {}

        config = self._get_config()
        output_dir = config.get("output_dir", "")
        watch_folder = config.get("watch_folder", "")

        # 从 Excel 读取
        excel_data = self._get_excel_data(self.current_type, self.current_year)
        for number, (source_file, meeting_date, pdf_path) in excel_data.items():
            self.meeting_data[number] = (source_file, meeting_date, pdf_path)

        # 从 PDF 文件夹补充
        pdf_data = self._scan_pdf_folder(self.current_type, self.current_year)
        for number, (source_file, meeting_date, pdf_path) in pdf_data.items():
            if number not in self.meeting_data:
                self.meeting_data[number] = (source_file, meeting_date, pdf_path)

        self._render_heatmap()

    def _get_excel_data(self, meeting_type: str, year: str) -> dict:
        """
        从 Excel 读取指定类型和年份的数据
        :return: {number: (source_file, meeting_date, pdf_path)} 按纪要序号索引
        """
        config = self._get_config()
        output_dir = config.get("output_dir", "")
        watch_folder = config.get("watch_folder", "")

        if not output_dir or not os.path.exists(output_dir):
            logger.warning(f"输出目录不存在: {output_dir}")
            return {}

        # 查找 Excel 文件
        excel_path = None
        for f in os.listdir(output_dir):
            if f.startswith("会议纪要汇总") and f.endswith(".xlsx"):
                excel_path = os.path.join(output_dir, f)
                break

        if not excel_path or not os.path.exists(excel_path):
            logger.warning(f"Excel 文件不存在: {output_dir}")
            return {}

        result = {}  # {number: (source_file, meeting_date, pdf_path)}
        seen_numbers = set()

        # Sheet名映射
        sheet_name_map = {
            "党委会": ["党委会", "party"],
            "董事会": ["董事会", "board"],
            "总经办": ["总经办", "gm"],
        }

        try:
            wb = load_workbook(excel_path, data_only=True)

            # 查找对应 Sheet
            target_sheets = sheet_name_map.get(meeting_type, [meeting_type])
            ws = None
            for s in wb.sheetnames:
                if s in target_sheets or meeting_type in s:
                    ws = wb[s]
                    break

            if ws is None:
                logger.warning(f"未找到会议类型 '{meeting_type}' 的 Sheet")
                wb.close()
                return {}

            # 读取数据（跳过表头行）
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue

                record_number = str(row[0] or "").strip()
                meeting_date = str(row[1] or "").strip()
                row_year = str(row[3] or "").strip()
                source_file = str(row[8] or "").strip()

                # 检查年份匹配
                if row_year != year:
                    continue

                # 从纪要编号提取数字
                number = self._extract_number(record_number)
                if number <= 0 or number > 50:
                    continue

                # 去重
                if number in seen_numbers:
                    continue
                seen_numbers.add(number)

                # 构建 PDF 路径（目录结构：监听目录/年份/会议类型/文件名）
                pdf_path = ""
                if watch_folder and source_file:
                    pdf_path = os.path.normpath(os.path.join(watch_folder, year, meeting_type, source_file))
                    # 如果文件不存在，尝试括号归一化匹配
                    if not os.path.exists(pdf_path):
                        found = self._find_pdf_with_brackets(
                            os.path.normpath(os.path.join(watch_folder, year, meeting_type)),
                            source_file
                        )
                        if found:
                            pdf_path = found

                result[number] = (source_file, meeting_date, pdf_path)

            wb.close()

        except Exception as e:
            logger.error(f"读取 Excel 数据失败: {e}")

        return result

    def _scan_pdf_folder(self, meeting_type: str, year: str) -> dict:
        """
        扫描监听目录下的 PDF 文件，补充没有 Excel 记录的文件
        :return: {number: (source_file, meeting_date, pdf_path)}
        """
        config = self._get_config()
        watch_folder = config.get("watch_folder", "")
        if not watch_folder:
            return {}

        year_dir = os.path.normpath(os.path.join(watch_folder, year, meeting_type))
        if not os.path.exists(year_dir):
            return {}

        result = {}
        # 匹配所有括号格式：〔〕【】﹝﹞（）() 「」
        bracket_open = r'[〔【﹝（(「『]'
        bracket_close = r'[〕】﹞）)」』]'
        pdf_pattern = re.compile(
            rf'{re.escape(meeting_type)}{bracket_open}{re.escape(year)}{bracket_close}(\d+)号\.pdf$'
        )

        for f in os.listdir(year_dir):
            match = pdf_pattern.match(f)
            if not match:
                continue

            number = int(match.group(1))
            if number < 1 or number > 50:
                continue

            source_file = f"{year}/{meeting_type}/{f}"
            pdf_path = os.path.join(year_dir, f)
            result[number] = (source_file, "", pdf_path)

        return result

    def _render_heatmap(self):
        """渲染热力图 - 5行×10列，编号1-50"""
        # 清除旧内容
        for widget in self.scrollable_frame.winfo_children():
            widget.destroy()

        # 颜色
        color = MEETING_COLORS.get(self.current_type, "#4A90D9")

        # 统计
        total_count = len(self.meeting_data)
        config = self._get_config()
        watch_folder = config.get("watch_folder", "")
        self.info_label.config(
            text=f"共 {total_count} 份纪要 | 监听目录: {watch_folder or '未配置'}"
        )

        # 渲染 5行 × 10列 共50个格子
        for row_idx in range(NUM_ROWS):
            for col_idx in range(COLS_PER_ROW):
                number = row_idx * COLS_PER_ROW + col_idx + 1  # 序号 1-50

                if number in self.meeting_data:
                    source_file, meeting_date, pdf_path = self.meeting_data[number]
                    # 有纪要 → 可点击按钮
                    btn = tk.Button(
                        self.scrollable_frame,
                        text=str(number),
                        font=("微软雅黑", 14, "bold"),
                        bg=color,
                        fg=TEXT_COLOR,
                        width=6,
                        height=2,
                        relief=tk.RAISED,
                        borderwidth=2,
                        cursor="hand2",
                        activebackground=self._lighten_color(color, 30),
                    )
                    btn.config(
                        command=lambda sf=source_file, num=number, pp=pdf_path:
                            self._open_file(sf, num, pp)
                    )
                    btn.grid(
                        row=row_idx, column=col_idx,
                        padx=GAP, pady=GAP, sticky="nsew"
                    )

                    # tooltip
                    self._create_tooltip(btn, self._format_tooltip(
                        number, meeting_type=self.current_type, date=meeting_date
                    ))
                else:
                    # 无纪要 → 灰色不可点击
                    lbl = tk.Label(
                        self.scrollable_frame,
                        text=str(number),
                        font=("微软雅黑", 14, "bold"),
                        bg=GRAY_COLOR,
                        fg=GRAY_TEXT,
                        width=6,
                        height=2,
                        relief=tk.SUNKEN,
                        borderwidth=1,
                    )
                    lbl.grid(
                        row=row_idx, column=col_idx,
                        padx=GAP, pady=GAP, sticky="nsew"
                    )

        # 配置列权重
        for col_idx in range(COLS_PER_ROW):
            self.scrollable_frame.columnconfigure(col_idx, weight=1)
        for row_idx in range(NUM_ROWS):
            self.scrollable_frame.rowconfigure(row_idx, weight=1)

    @staticmethod
    def _lighten_color(hex_color: str, amount: int) -> str:
        """使颜色变亮"""
        hex_color = hex_color.lstrip("#")
        r, g, b = int(hex_color[:2], 16), int(hex_color[2:4], 16), int(hex_color[4:], 16)
        r = min(255, r + amount)
        g = min(255, g + amount)
        b = min(255, b + amount)
        return f"#{r:02x}{g:02x}{b:02x}"

    def _format_tooltip(self, number: int, meeting_type: str = "", date: str = "") -> str:
        """格式化 tooltip 文本"""
        lines = [f"第 {number} 号"]
        if meeting_type:
            lines.insert(0, meeting_type)
        if date:
            lines.append(f"日期: {date}")
        return "\n".join(lines)

    def _create_tooltip(self, widget, text: str):
        """为组件创建 tooltip"""
        tooltip = None

        def show(event):
            nonlocal tooltip
            if tooltip:
                return
            x = event.x_root + 15
            y = event.y_root + 10
            tooltip = tk.Toplevel(widget)
            tooltip.wm_overrideredirect(True)
            tooltip.wm_geometry(f"+{x}+{y}")
            lbl = tk.Label(
                tooltip, text=text, justify=tk.LEFT,
                bg="#FFFFDD", fg="#333333",
                font=("微软雅黑", 9),
                padx=6, pady=3,
                relief=tk.SOLID, borderwidth=1
            )
            lbl.pack()

        def hide(event):
            nonlocal tooltip
            if tooltip:
                tooltip.destroy()
                tooltip = None

        widget.bind("<Enter>", show)
        widget.bind("<Leave>", hide)

    def _open_file(self, source_file: str, number: int, pdf_path: str = ""):
        """打开纪要文件"""
        # 优先使用传入的 pdf_path
        if not pdf_path:
            config = self._get_config()
            watch_folder = config.get("watch_folder", "")
            if watch_folder and source_file:
                # source_file 可能是纯文件名也可能是相对路径，统一处理
                pdf_path = os.path.normpath(os.path.join(watch_folder, source_file))
                # 如果文件不存在，尝试按目录结构拼接（年份/会议类型/文件名）
                if not os.path.exists(pdf_path) and "/" not in source_file:
                    pdf_path = os.path.normpath(os.path.join(
                        watch_folder, self.current_year, self.current_type, source_file
                    ))

        # 如果路径不存在，用智能括号匹配查找
        if pdf_path and not os.path.exists(pdf_path) and watch_folder:
            # 先尝试在当前目录查找
            found = self._find_pdf_with_brackets(os.path.dirname(pdf_path), os.path.basename(pdf_path))
            if found:
                pdf_path = found
            elif "/" not in source_file:
                # 再尝试在年份/类型目录下查找
                year_dir = os.path.normpath(os.path.join(watch_folder, self.current_year, self.current_type))
                found = self._find_pdf_with_brackets(year_dir, source_file)
                if found:
                    pdf_path = found

        if pdf_path and os.path.exists(pdf_path):
            try:
                if os.name == "nt":
                    os.startfile(pdf_path)
                elif sys.platform == "darwin":
                    subprocess.run(["open", pdf_path], check=True)
                else:
                    subprocess.run(["xdg-open", pdf_path], check=True)
                logger.info(f"已打开文件: {pdf_path}")
            except Exception as e:
                logger.error(f"打开文件失败: {e}")
                messagebox.showerror(
                    "打开失败",
                    f"无法打开文件：\n{pdf_path}\n\n请检查系统默认 PDF 阅读器"
                )
        else:
            messagebox.showinfo(
                "文件不存在",
                f"第 {number} 号纪要文件不存在：\n{source_file}\n\n请确认文件是否在监听目录中。"
            )

    @staticmethod
    def _normalize_brackets(filename: str, to_square: bool = False) -> str:
        """
        归一化文件名中的括号，将所有括号统一为一种格式
        :param filename: 原始文件名
        :param to_square: True → 转成【】；False → 转成〔〕（默认）
        处理的括号对：
          【】〔〕﹝﹞  （中文专用括号）
          「」『』  （中文引号类括号）
          （）()   （全角/半角圆括号）
        """
        # 所有右括号统一
        result = filename

        if to_square:
            # 全部转成 【】
            result = result.replace("〔", "【").replace("〕", "】")
            result = result.replace("﹝", "【").replace("﹞", "】")
            result = result.replace("「", "【").replace("」", "】")
            result = result.replace("『", "【").replace("』", "】")
            result = result.replace("（", "【").replace("）", "】")
            result = result.replace("(", "【").replace(")", "】")
        else:
            # 全部转成 〔〕（默认）
            result = result.replace("【", "〔").replace("】", "〕")
            result = result.replace("﹝", "〔").replace("﹞", "〕")
            result = result.replace("「", "〔").replace("」", "〕")
            result = result.replace("『", "〔").replace("』", "〕")
            result = result.replace("（", "〔").replace("）", "〕")
            result = result.replace("(", "〔").replace(")", "〕")

        return result

    @staticmethod
    def _find_pdf_with_brackets(directory: str, base_name: str) -> str:
        """
        在指定目录中查找 PDF 文件，自动处理括号格式不匹配
        例如目录中有 总经办〔2026〕3号.pdf，但 base_name 是 总经办【2026】3号.pdf
        :return: 实际文件路径，未找到返回空字符串
        """
        if not base_name or not directory or not os.path.isdir(directory):
            return ""

        # 先尝试精确匹配原始文件名
        exact_path = os.path.join(directory, base_name)
        if os.path.exists(exact_path):
            return exact_path

        # 尝试括号互换后的文件名
        alt_name = HeatmapView._normalize_brackets(base_name)
        alt_path = os.path.join(directory, alt_name)
        if os.path.exists(alt_path):
            return alt_path

        # 最后扫描目录，做一次宽松匹配（忽略括号差异）
        for f in os.listdir(directory):
            if not f.lower().endswith('.pdf'):
                continue
            # 归一化后再比较
            norm_f = HeatmapView._normalize_brackets(f)
            norm_base = HeatmapView._normalize_brackets(base_name)
            if norm_f == norm_base:
                return os.path.join(directory, f)

        return ""

    @staticmethod
    def _extract_number(record_number: str) -> int:
        """从纪要编号提取数字"""
        # 匹配所有括号格式：〔〕【】﹝﹞（）()「」
        bracket_open = r'[〔【﹝（(「『]'
        bracket_close = r'[〕】﹞）)」』]'
        match = re.search(rf'{bracket_open}\d+{bracket_close}(\d+)号', record_number)
        if match:
            return int(match.group(1))
        match = re.search(r'(\d+)号', record_number)
        if match:
            return int(match.group(1))
        return 0
